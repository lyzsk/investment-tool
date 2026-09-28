# compare_equity.py — A/C/D 三组 replay 对比产物(2026-09-27 用户定呈现方式)
# 输入: <run-dir>/<group>/{equity.csv, trades.csv}  (replay/settle.mjs 口径, 带 BOM)
# 输出: <run-dir>/compare/equity-compare.png   顶部三组收益叠加(右上证) + 每组独立子图标注买卖票名
#       <run-dir>/compare/daily_detail.xlsx    每日明细: 三组累计/日收益 + 每组当日买/卖票名@价×量
# 用法: C:/Users/admin/anaconda3/python.exe compare_equity.py [--run-dir runs/replay-260629_260924]
#       组目录缺失只警告不崩(跑批中途也能出半成品); A 组旧目录可用 --dir A=<path> 覆盖
import argparse
import csv
import os
from datetime import datetime

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei']
plt.rcParams['axes.unicode_minus'] = False

BASE = os.path.dirname(os.path.abspath(__file__))
GROUPS = [('A', 'A·照抄桃哥', '#1f77b4'), ('C', 'C·裸LLM', '#ff7f0e'), ('D', 'D·persona', '#d62728')]
CN_RED, CN_GREEN = '#d62728', '#2ca02c'  # A股配色: 红买/涨, 绿卖/跌


def read_csv(path):
    with open(path, encoding='utf-8-sig') as f:  # utf-8-sig 自动吃 BOM
        return list(csv.DictReader(f))


def load_group(gdir):
    eq = read_csv(os.path.join(gdir, 'equity.csv'))
    dates = [datetime.strptime(r['date'], '%Y%m%d') for r in eq]
    cum = [float(r['return_pct']) for r in eq]
    sh = [float(r['sh_close']) if r.get('sh_close') else None for r in eq]
    trs = [r for r in read_csv(os.path.join(gdir, 'trades.csv')) if r['verdict'] == 'filled']
    # 每日操作: 成交日(fill_bar 前8位)聚合 "name@px×qty"
    ops = {}
    for t in trs:
        d = (t.get('fill_bar') or t['trigger_bar'])[:8]
        ops.setdefault(d, {'buy': [], 'sell': []})[t['side']].append(
            f"{t['name']}@{float(t['px']):.2f}×{t['qty']}")
    return {'dates': dates, 'cum': cum, 'sh': sh, 'ops': ops, 'trs': trs}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--run-dir', default=os.path.join(BASE, 'runs', 'replay-260629_260924'))
    ap.add_argument('--dir', action='append', default=[], help='组目录覆盖, 形如 A=<path>(可多次)')
    args = ap.parse_args()
    overrides = dict(x.split('=', 1) for x in args.dir)
    out_dir = os.path.join(args.run_dir, 'compare')
    os.makedirs(out_dir, exist_ok=True)

    data = {}
    for g, label, _ in GROUPS:
        gdir = overrides.get(g, os.path.join(args.run_dir, g))
        if os.path.exists(os.path.join(gdir, 'equity.csv')):
            data[g] = load_group(gdir)
            print(f"[{g}] {len(data[g]['dates'])} 天, {len(data[g]['trs'])} 笔成交, "
                  f"终值 {data[g]['cum'][-1]:+.2f}%")
        else:
            print(f"[{g}] 缺 {gdir}/equity.csv, 跳过")

    if not data:
        raise SystemExit("没有任何一组数据, 检查 --run-dir")

    # ---------- 图: 顶部叠加 + 每组独立子图 ----------
    n = len(data)
    fig, axes = plt.subplots(1 + n, 1, figsize=(17, 5 + 4 * n), sharex=True,
                             gridspec_kw={'height_ratios': [1.3] + [1] * n})
    ax0 = axes[0]
    for g, label, color in GROUPS:
        if g in data:
            d = data[g]
            ax0.plot(d['dates'], d['cum'], color=color, lw=2,
                     label=f"{label} ({d['cum'][-1]:+.2f}%)")
    sh_src = next((data[g] for g, _, _ in GROUPS if g in data and any(data[g]['sh'])), None)
    if sh_src:
        pairs = [(d, s) for d, s in zip(sh_src['dates'], sh_src['sh']) if s]
        ax0b = ax0.twinx()
        ax0b.plot([p[0] for p in pairs], [p[1] for p in pairs], color='#888888',
                  lw=1, ls=':', label='上证指数(右轴)')
        ax0b.tick_params(axis='y', colors='#888888')
    ax0.axhline(0, color='#cccccc', lw=0.8)
    ax0.legend(loc='upper left')
    ax0.set_title('A/C/D 三组收益率对比(replay 回测, 次根m30开盘价成交)')
    ax0.set_ylabel('累计收益 %')
    ax0.grid(alpha=0.3)

    row = 1
    for g, label, color in GROUPS:
        if g not in data:
            continue
        ax, d = axes[row], data[g]
        row += 1
        ax.plot(d['dates'], d['cum'], color=color, lw=1.5)
        ax.axhline(0, color='#cccccc', lw=0.8)
        ax.set_ylabel(f'{label} %', fontsize=9)
        ax.grid(alpha=0.3)
        # y 轴上下各留 20% 头部空间, 标注文字才不会飘出子图(9/27 实测 sell 标注飘进面板间隙)
        lo, hi = min(d['cum']), max(d['cum'])
        pad = max((hi - lo) * 0.2, 0.5)
        ax.set_ylim(lo - pad, hi + pad)
        ax.margins(x=0.02)
        ds = {dt.strftime('%Y%m%d'): dt for dt in d['dates']}
        cum_by = {dt.strftime('%Y%m%d'): c for dt, c in zip(d['dates'], d['cum'])}
        i = 0
        for day in sorted(d['ops']):
            if day not in ds:
                continue
            x, y = ds[day], cum_by[day]
            for side, marker, c, dy in (('buy', '^', CN_RED, -1), ('sell', 'v', CN_GREEN, 1)):
                names = [s.split('@')[0] for s in d['ops'][day][side]]
                if not names:
                    continue
                ax.scatter([x], [y], marker=marker, color=c, s=45, zorder=5)
                ax.annotate('/'.join(names), (x, y), textcoords='offset points',
                            xytext=(6 * ((i % 3) - 1), (8 + 7 * (i % 2)) * dy),
                            rotation=40, fontsize=6.5, color=c, ha='center',
                            annotation_clip=False)
                i += 1
    axes[-1].xaxis.set_major_formatter(mdates.DateFormatter('%m-%d'))
    axes[-1].xaxis.set_major_locator(mdates.WeekdayLocator(byweekday=0))
    fig.tight_layout()
    png = os.path.join(out_dir, 'equity-compare.png')
    fig.savefig(png, dpi=140)
    print(f"图: {png}")

    # ---------- Excel 每日明细 ----------
    wb = Workbook()
    ws = wb.active
    ws.title = '每日明细'
    head = ['日期', '星期', '上证收盘',
            'A累计%', 'C累计%', 'D累计%', 'A日%', 'C日%', 'D日%',
            'A买入', 'A卖出', 'C买入', 'C卖出', 'D买入', 'D卖出']
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDEBF7')
        c.alignment = Alignment(horizontal='center')
    all_days = sorted({dt.strftime('%Y%m%d') for g in data for dt in data[g]['dates']})
    cum_map = {g: {dt.strftime('%Y%m%d'): c for dt, c in zip(data[g]['dates'], data[g]['cum'])} for g in data}
    sh_map = {}
    for g in data:
        for dt, s in zip(data[g]['dates'], data[g]['sh']):
            if s:
                sh_map[dt.strftime('%Y%m%d')] = s
    prev = {g: 0.0 for g in data}
    wk = '一二三四五六日'
    for day in all_days:
        dt = datetime.strptime(day, '%Y%m%d')
        row = [dt.strftime('%Y-%m-%d'), wk[dt.weekday()], sh_map.get(day, '')]
        for g, _, _ in GROUPS:
            c = cum_map.get(g, {}).get(day)
            row.append(round(c, 2) if c is not None else '')
        for g, _, _ in GROUPS:
            c = cum_map.get(g, {}).get(day)
            row.append(round(c - prev[g], 2) if c is not None else '')
            if c is not None:
                prev[g] = c
        for g, _, _ in GROUPS:
            ops = data.get(g, {}).get('ops', {}).get(day, {})
            row.append('; '.join(ops.get('buy', [])))
            row.append('; '.join(ops.get('sell', [])))
        ws.append(row)
    widths = [11, 6, 10, 9, 9, 9, 8, 8, 8, 28, 28, 28, 28, 28, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'D2'
    for r in ws.iter_rows(min_row=2):
        for idx in (3, 4, 5, 6, 7, 8):  # 收益列红涨绿跌
            v = r[idx].value
            if isinstance(v, (int, float)):
                r[idx].font = Font(color=(CN_RED if v > 0 else CN_GREEN if v < 0 else '000000').lstrip('#'))
    # 交易全量 sheet
    ws2 = wb.create_sheet('成交全量')
    ws2.append(['组', '日期', '代码', '名称', '方向', '成交价', '量', '费用', '备注'])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for g, _, _ in GROUPS:
        for t in data.get(g, {}).get('trs', []):
            ws2.append([g, (t.get('fill_bar') or t['trigger_bar'])[:8], t['code'], t['name'],
                        '买' if t['side'] == 'buy' else '卖', float(t['px']), int(t['qty']),
                        round(float(t['fee']), 2), t.get('note', '')])
    for i, w in enumerate([4, 10, 9, 10, 5, 8, 8, 8, 40], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    xlsx = os.path.join(out_dir, 'daily_detail.xlsx')
    wb.save(xlsx)
    print(f"表: {xlsx}")


if __name__ == '__main__':
    main()
